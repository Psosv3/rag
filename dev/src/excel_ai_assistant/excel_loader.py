# excel_ai_assistant/excel_loader.py
from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .exceptions import ExcelLoadError
from .models import CellInfo, SheetSummary, WorkbookSummary, LabeledValue

logger = logging.getLogger(__name__)


@dataclass
class ExcelWorkbookLoader:
    """
    Load and inspect Excel workbooks (.xlsx) to produce structured summaries.

    This loader uses two views of the same workbook:
      - formulas_wb: data_only=False (formulas are visible)
      - values_wb:   data_only=True  (cached formula results are visible)

    It is stateless and safe to use across threads/processes.
    """

    def load_summary(
        self,
        source: Union[str, Path, bytes],
        file_id: Optional[str] = None,
    ) -> WorkbookSummary:
        """
        Load an Excel file and build a WorkbookSummary.

        :param source: Path to .xlsx file or raw bytes.
        :param file_id: Logical identifier for the file (e.g., UUID). If None, uses path or "in-memory".
        """
        try:
            if isinstance(source, (str, Path)):
                path = Path(source)
                formulas_wb = load_workbook(filename=path, data_only=False, read_only=False)
                values_wb = load_workbook(filename=path, data_only=True, read_only=False)
                logical_id = file_id or path.name
            else:
                from io import BytesIO

                bytes_data = bytes(source)
                formulas_wb = load_workbook(filename=BytesIO(bytes_data), data_only=False, read_only=False)
                values_wb = load_workbook(filename=BytesIO(bytes_data), data_only=True, read_only=False)
                logical_id = file_id or "in-memory-workbook"
        except Exception as exc:
            logger.exception("Failed to load workbook")
            raise ExcelLoadError(f"Failed to load Excel workbook: {exc}") from exc

        try:
            formula_sheets_by_title: Dict[str, Worksheet] = {
                ws.title: ws for ws in formulas_wb.worksheets
            }
            value_sheets_by_title: Dict[str, Worksheet] = {
                ws.title: ws for ws in values_wb.worksheets
            }
        except Exception as exc:
            logger.exception("Failed to index worksheets by title")
            raise ExcelLoadError(f"Failed to index worksheets: {exc}") from exc

        sheets: List[SheetSummary] = []

        for sheet_title, ws_formulas in formula_sheets_by_title.items():
            ws_values = value_sheets_by_title.get(sheet_title)
            if ws_values is None:
                logger.warning(
                    "Sheet %s present in formulas workbook but not in values workbook; skipping.",
                    sheet_title,
                )
                continue

            try:
                sheet_summary = self._summarize_sheet(
                    ws_formulas=ws_formulas,
                    ws_values=ws_values,
                )
                sheets.append(sheet_summary)
            except Exception as exc:
                logger.exception("Failed to summarize sheet %s", sheet_title)
                raise ExcelLoadError(f"Failed to summarize sheet '{sheet_title}': {exc}") from exc

        return WorkbookSummary(file_id=logical_id, sheets=sheets)

    # -------------------------------------------------------------------------
    # Sheet-level summarization
    # -------------------------------------------------------------------------

    def _summarize_sheet(
        self,
        ws_formulas: Worksheet,
        ws_values: Worksheet,
    ) -> SheetSummary:
        # Sizes should match between formula and value workbooks
        n_rows = ws_values.max_row or ws_formulas.max_row or 0
        n_cols = ws_values.max_column or ws_formulas.max_column or 0

        headers, header_row_idx = self._infer_headers(ws_values)
        # Sample rows starting from A1 (row 1, column 1) so that indices always
        # correspond to the global Excel grid (Row = Excel row, ColumnN = Excel column N).
        sample_rows = self._sample_rows_grid(
            ws=ws_values,
            max_rows=n_rows,
        )
        formula_cells, cross_sheet_refs = self._collect_formulas(ws_formulas, ws_values)
        labeled_values = self._extract_labeled_values(ws_values)

        return SheetSummary(
            name=ws_formulas.title,
            n_rows=n_rows,
            n_cols=n_cols,
            headers=headers,
            sample_rows=sample_rows,
            formula_cells=formula_cells,
            cross_sheet_references=cross_sheet_refs,
            labeled_values=labeled_values,
        )

    # -------------------------------------------------------------------------
    # Header inference (using values workbook)
    # -------------------------------------------------------------------------

    def _infer_headers(self, ws: Worksheet) -> tuple[List[str], int]:
        """
        Heuristic: take the first non-empty row with at least 2 non-empty cells as header row.

        Returns (headers, header_row_index). If no headers found, returns ([], 0).

        Uses values_only=True, so each row is a tuple of raw values (including formula results).
        """
        max_row = ws.max_row or 0
        if max_row == 0:
            return [], 0

        scan_max_row = min(max_row, 10)
        header_row_idx = 0
        headers: List[str] = []

        for row_idx, row in enumerate(
            ws.iter_rows(min_row=1, max_row=scan_max_row, values_only=True),
            start=1,
        ):
            values = list(row)
            if not values:
                continue

            non_empty = [
                v for v in values
                if v is not None and str(v).strip() != ""
            ]
            if len(non_empty) >= 2:
                header_row_idx = row_idx
                headers = [str(v) if v is not None else "" for v in values]
                break

        return headers, header_row_idx

    # -------------------------------------------------------------------------
    # Row sampling (grid-style, using Column1..ColumnN, starting at A1)
    # -------------------------------------------------------------------------

    def _sample_rows_grid(
        self,
        ws: Worksheet,
        max_rows: int,
    ) -> List[Dict[str, Any]]:
        """
        Sample rows as a raw grid, starting from cell A1.

        Keys are always Column1..ColumnN, where N is the sheet's max_column.

        Row and column indices are therefore aligned with the Excel grid:
        - Column1 corresponds to column A, Column2 to B, etc.
        - The first dictionary in the list represents Excel row 1, the second row 2, etc.
        """
        if max_rows <= 0:
            return []

        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        if max_row == 0 or max_col == 0:
            return []

        start_row = 1
        end_row = min(max_row, max_rows)

        samples: List[Dict[str, Any]] = []
        for row_values in ws.iter_rows(
            min_row=start_row,
            max_row=end_row,
            min_col=1,
            max_col=max_col,
            values_only=True,
        ):
            row_dict: Dict[str, Any] = {}
            for col_idx, cell_value in enumerate(row_values, start=1):
                col_name = f"Column{col_idx}"
                row_dict[col_name] = cell_value
            samples.append(row_dict)

        return samples

    # -------------------------------------------------------------------------
    # Formula and cross-sheet reference collection
    # -------------------------------------------------------------------------

    def _collect_formulas(
        self,
        ws_formulas: Worksheet,
        ws_values: Worksheet,
    ) -> tuple[List[CellInfo], List[str]]:
        """
        Collect cells that contain formulas and detect cross-sheet references.

        For each formula cell, we also fetch the cached value from the values workbook.
        """
        formula_cells: List[CellInfo] = []
        cross_sheet_refs: List[str] = []

        sheet_ref_pattern = re.compile(r"(?P<sheet>'[^']+'|[A-Za-z0-9_]+)!")

        max_row = ws_formulas.max_row or 0
        max_col = ws_formulas.max_column or 0

        for row_idx in range(1, max_row + 1):
            for col_idx in range(1, max_col + 1):
                f_cell = ws_formulas.cell(row=row_idx, column=col_idx)
                value_cell = ws_values.cell(row=row_idx, column=col_idx)

                f_value = f_cell.value
                if isinstance(f_value, str) and f_value.startswith("="):
                    addr = f_cell.coordinate
                    cached_val = value_cell.value

                    formula_cells.append(
                        CellInfo(
                            sheet_name=ws_formulas.title,
                            address=addr,
                            value=cached_val,
                            formula=f_value,
                        )
                    )

                    for match in sheet_ref_pattern.finditer(f_value):
                        sheet_name = match.group("sheet").strip("'")
                        if sheet_name != ws_formulas.title:
                            cross_sheet_refs.append(sheet_name)

        return formula_cells, cross_sheet_refs

    # -------------------------------------------------------------------------
    # Labeled value extraction (generic label -> value semantics, using values workbook)
    # -------------------------------------------------------------------------

    def _extract_labeled_values(self, ws: Worksheet) -> List[LabeledValue]:
        """
        Detect label->value patterns, both horizontal and vertical, using the values workbook.

        Examples (generic, not specific to any domain):
          B1: "Chiffre d'affaire année 1"  C1: 324654532
          A5: "Nombre d'employés"         B5: 45
          D10: "CPU usage (max)"          E10: 0.87

        Heuristics:
          - label cell: string with at least one letter
          - adjacent value cell: numeric / date / bool / numeric-like string
        """

        def is_potential_label(v: Any) -> bool:
            if not isinstance(v, str):
                return False
            s = v.strip()
            if not s:
                return False
            return any(ch.isalpha() for ch in s)

        def is_potential_value(v: Any) -> bool:
            from datetime import date, datetime

            if v is None:
                return False
            if isinstance(v, (int, float, bool, date, datetime)):
                return True
            if isinstance(v, str):
                s = v.strip()
                if not s:
                    return False
                compact = s.replace(" ", "").replace(",", "").replace(".", "")
                return compact.isdigit()
            return False

        labeled_values: List[LabeledValue] = []

        max_row = ws.max_row or 0
        max_col = ws.max_column or 0

        if max_row == 0 or max_col == 0:
            return []

        # Horizontal: label in col j, value in col j+1 (same row)
        for row_idx in range(1, max_row + 1):
            for col_idx in range(1, max_col):
                label_cell = ws.cell(row=row_idx, column=col_idx)
                value_cell = ws.cell(row=row_idx, column=col_idx + 1)

                label_val = label_cell.value
                value_val = value_cell.value

                if is_potential_label(label_val) and is_potential_value(value_val):
                    labeled_values.append(
                        LabeledValue(
                            sheet_name=ws.title,
                            label=str(label_val).strip(),
                            label_address=label_cell.coordinate,
                            value_address=value_cell.coordinate,
                            value=value_val,
                            orientation="horizontal",
                        )
                    )

        # Vertical: label in row i, value in row i+1 (same column)
        for col_idx in range(1, max_col + 1):
            for row_idx in range(1, max_row):
                label_cell = ws.cell(row=row_idx, column=col_idx)
                value_cell = ws.cell(row=row_idx + 1, column=col_idx)

                label_val = label_cell.value
                value_val = value_cell.value

                if is_potential_label(label_val) and is_potential_value(value_val):
                    labeled_values.append(
                        LabeledValue(
                            sheet_name=ws.title,
                            label=str(label_val).strip(),
                            label_address=label_cell.coordinate,
                            value_address=value_cell.coordinate,
                            value=value_val,
                            orientation="vertical",
                        )
                    )

        # Deduplicate
        seen = set()
        unique_labeled_values: List[LabeledValue] = []
        for lv in labeled_values:
            key = (lv.sheet_name, lv.label_address, lv.value_address)
            if key in seen:
                continue
            seen.add(key)
            unique_labeled_values.append(lv)

        return unique_labeled_values
